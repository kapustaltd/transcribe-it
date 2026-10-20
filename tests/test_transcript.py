import io
import unittest

from docx import Document

from transcript import Line, Section, Transcript, merge, render_docx, split_chunks


def sample(title="Транскрипция: Тест", participants=None, sections=None) -> Transcript:
    return Transcript(
        title=title,
        description="Описание",
        participants=participants if participants is not None else ["Каша (Вова)", "Азазин (Рома)"],
        kind="Разговорный стрим-подкаст",
        sections=sections if sections is not None else [
            Section(title="Приветствие", lines=[
                Line(speaker="Каша", text="Всем привет!"),
                Line(speaker="Азазин", text="Здравствуйте."),
            ]),
            Section(title="Пустой раздел", lines=[]),
            Section(title="Монолог", lines=[Line(speaker="", text="Говорю один.")]),
        ],
    )


class SplitChunksTest(unittest.TestCase):
    def test_keeps_lines_whole_and_respects_limit(self):
        text = "\n".join(["a" * 40] * 10)
        chunks = split_chunks(text, limit=100)
        self.assertTrue(all(len(chunk) <= 100 for chunk in chunks))
        self.assertEqual("\n".join(chunks), text)

    def test_short_text_is_one_chunk(self):
        self.assertEqual(split_chunks("короткий текст"), ["короткий текст"])


class MergeTest(unittest.TestCase):
    def test_header_comes_from_first_part_and_sections_are_concatenated(self):
        first = sample(title="Первый", sections=[Section(title="Один", lines=[])])
        second = sample(title="", participants=[], sections=[Section(title="Два", lines=[])])
        merged = merge([first, second])
        self.assertEqual(merged.title, "Первый")
        self.assertEqual([s.title for s in merged.sections], ["Один", "Два"])


class RenderDocxTest(unittest.TestCase):
    def read(self, transcript: Transcript) -> Document:
        return Document(io.BytesIO(render_docx(transcript)))

    def test_headings_are_numbered_and_empty_sections_skipped(self):
        headings = [p.text for p in self.read(sample()).paragraphs if p.style.name == "Heading 1"]
        self.assertEqual(headings, ["Часть 1. Приветствие", "Часть 2. Монолог"])

    def test_speaker_names_are_bold_and_text_is_plain(self):
        paragraphs = [p for p in self.read(sample()).paragraphs if p.text.startswith("Каша:")]
        self.assertEqual(len(paragraphs), 1)
        speaker, text = paragraphs[0].runs[0], paragraphs[0].runs[1]
        self.assertEqual(speaker.text, "Каша: ")
        self.assertTrue(speaker.bold)
        self.assertFalse(text.bold)
        self.assertEqual(text.text, "Всем привет!")

    def test_monologue_has_no_speaker_prefix(self):
        texts = [p.text for p in self.read(sample()).paragraphs]
        self.assertIn("Говорю один.", texts)

    def test_title_and_meta_are_rendered(self):
        document = self.read(sample())
        self.assertEqual(document.paragraphs[0].text, "Транскрипция: Тест")
        meta = next(p.text for p in document.paragraphs if p.text.startswith("Участники:"))
        self.assertIn("Каша (Вова), Азазин (Рома)", meta)
        self.assertIn("Формат: Разговорный стрим-подкаст", meta)


if __name__ == "__main__":
    unittest.main()
